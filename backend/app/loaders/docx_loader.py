"""DOCX document loader using python-docx."""

from pathlib import Path
import docx
from fastapi import HTTPException, status

from app.loaders.base_loader import BaseLoader
from app.schemas.document import ExtractedDocument


class DocxLoader(BaseLoader):
    """Extracts paragraphs and tables from Word (.docx) documents."""

    def supported_extensions(self) -> list[str]:
        return [".docx"]

    def extract(self, file_path: Path) -> ExtractedDocument:
        if not file_path.exists():
            raise FileNotFoundError(f"El archivo {file_path} no existe.")

        try:
            doc = docx.Document(str(file_path))
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"No se pudo leer el archivo Word .docx: {exc}",
            )

        content_parts: list[str] = []

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                content_parts.append(text)

        # Extract tables
        table_count = len(doc.tables)
        for t_idx, table in enumerate(doc.tables):
            table_lines: list[str] = [f"[Tabla {t_idx + 1}]"]
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                # Filter duplicate merged cells
                cleaned_cells = []
                for cell in cells:
                    if not cleaned_cells or cell != cleaned_cells[-1]:
                        cleaned_cells.append(cell)
                if any(cleaned_cells):
                    table_lines.append(" | ".join(cleaned_cells))
            if len(table_lines) > 1:
                content_parts.append("\n".join(table_lines))

        full_content = "\n\n".join(content_parts)
        if not full_content.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="El archivo .docx no contiene texto extraíble.",
            )

        # Document core properties
        doc_props = {}
        try:
            props = doc.core_properties
            doc_props = {
                "author": props.author or "",
                "title": props.title or "",
                "created": str(props.created) if props.created else "",
                "modified": str(props.modified) if props.modified else "",
                "paragraphs_count": len(doc.paragraphs),
                "tables_count": table_count,
            }
        except Exception:
            pass

        return ExtractedDocument(
            filename=file_path.name,
            file_type="docx",
            file_size_bytes=file_path.stat().st_size,
            content=full_content,
            metadata=doc_props,
        )
